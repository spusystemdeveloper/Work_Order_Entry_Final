import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\ADMIN\source\repos\Queries\ENLIVO_AND_AMBER_CONSOLIDATED.xlsx', read_only=True)
print("Sheet names:", wb.sheetnames)

sheet = wb[wb.sheetnames[0]]
rows = list(sheet.iter_rows(max_row=10, values_only=True))
for r in rows:
    print(r)
